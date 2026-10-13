"""Utility for extracting text from uploaded resume documents (PDF, DOCX, TXT)."""

import io
from typing import Optional
from pypdf import PdfReader
from docx import Document


def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """Extract plain text from an uploaded resume file based on its extension.

    Args:
        file_bytes: The raw bytes of the uploaded file.
        filename: Name of the uploaded file including extension.

    Returns:
        Extracted plain text content.

    Raises:
        ValueError: If the file type is unsupported or extraction yields no text.
    """
    ext = filename.lower().split(".")[-1] if "." in filename else ""

    if ext in ["txt", "text", "md"]:
        return file_bytes.decode("utf-8", errors="ignore").strip()

    elif ext == "pdf":
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            pages_text = []
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    pages_text.append(page_text)
            extracted = "\n".join(pages_text).strip()
            if not extracted:
                raise ValueError("PDF appears to be empty or contains scanned images without OCR text.")
            return extracted
        except Exception as e:
            raise ValueError(f"Failed to extract text from PDF: {str(e)}") from e

    elif ext in ["docx", "doc"]:
        try:
            doc = Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            extracted = "\n".join(paragraphs).strip()
            if not extracted:
                raise ValueError("Word document appears to be empty.")
            return extracted
        except Exception as e:
            raise ValueError(f"Failed to extract text from Word document: {str(e)}") from e

    else:
        raise ValueError(
            f"Unsupported file format '.{ext}'. Supported formats are: PDF (.pdf), Word (.docx), and Text (.txt)."
        )
